import csv
import io
from datetime import UTC, datetime
from uuid import uuid4
from xml.etree.ElementTree import ParseError, fromstring

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db.models import AuditLog, User
from app.domain import configuration_xml
from app.domain.exports import csv_cell, report_csv, report_xlsx, report_xml
from app.domain.reports import lesson_report
from app.domain.runtime_configuration import RuntimeConfiguration, read_snapshot
from tests.test_auth import sign_in
from tests.test_lessons import lesson_fixture


@pytest.mark.parametrize("value", ["=1+1", "+SUM(A1)", "-1+2", "@SUM(A1)", " \t=1+1"])
def test_csv_formula_protection(value: str) -> None:
    assert csv_cell(value) == "'" + value
    assert csv_cell(-3) == -3
    assert csv_cell(None) == ""


async def test_report_formats_preserve_columns_unicode_nulls_and_manual_mark(
    db: AsyncSession,
) -> None:
    lesson, rows = await lesson_fixture(db, 1)
    rows[0].state = "CLOSED"
    rows[0].delivered_at = datetime.now(UTC)
    rows[0].score = {"total": 80, "axes": {"status": {"score": 80, "weight": 1}}, "violations": []}
    rows[0].teacher_override = {"total": 95, "comment": "Учтён доклад"}
    await db.flush()
    data = await lesson_report(db, lesson)
    data["students"][0]["short_name"] = "=Учебный; <Иванов>"
    encoded = report_csv(data)
    assert encoded.startswith(b"\xef\xbb\xbf")
    parsed = list(csv.reader(io.StringIO(encoded.decode("utf-8-sig")), delimiter=";"))
    assert parsed[0] == data["columns"]
    assert parsed[1][0] == "'=Учебный; <Иванов>" and parsed[1][-1] == "95.0*"
    xml = fromstring(report_xml(data))
    assert xml.tag == "lesson_report" and xml.attrib == {"schema": "1"}
    assert xml.findtext("students/item/short_name") == "=Учебный; <Иванов>"
    assert xml.findtext("students/item/manually_corrected") == "true"
    assert xml.find("lesson/finished_at").attrib == {"nil": "true"}  # type: ignore[union-attr]


@pytest.mark.parametrize("value", ["=1+1", "+SUM(A1)", "-1+2", "@SUM(A1)", " \t=1+1"])
async def test_xlsx_preserves_text_without_formulas_and_numeric_cells(
    db: AsyncSession, value: str
) -> None:
    lesson, _ = await lesson_fixture(db, 1)
    data = await lesson_report(db, lesson)
    data["students"][0].update(short_name=value, total=92.5, manually_corrected=False)
    workbook = load_workbook(io.BytesIO(report_xlsx(data)))
    sheet = workbook.active
    assert sheet is not None
    assert [cell.value for cell in sheet[1]] == data["columns"]
    assert sheet["A2"].value == value and sheet["A2"].data_type == "s"
    assert sheet["G2"].value == 92.5 and sheet["G2"].data_type == "n"
    assert sheet["C2"].value is None
    data["students"][0]["manually_corrected"] = True
    corrected = load_workbook(io.BytesIO(report_xlsx(data))).active
    assert corrected is not None and corrected["G2"].value == "92.5*"


@pytest.mark.parametrize("extension", ["csv", "xml", "xlsx"])
async def test_report_export_is_owned_audited_and_private(
    db: AsyncSession, client: AsyncClient, extension: str
) -> None:
    lesson, _ = await lesson_fixture(db)
    path = f"/api/v1/teacher/lessons/{lesson.id}/report.{extension}"
    await sign_in(client)
    assert (await client.get(path)).status_code == 403
    await client.post("/api/v1/auth/logout")
    await sign_in(client, "teacher", "teacher")
    response = await client.get(path)
    assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
    assert await db.scalar(
        select(AuditLog.id).where(
            AuditLog.action == "LESSON_REPORT_EXPORTED",
            AuditLog.entity_id == lesson.id,
            AuditLog.payload["format"].astext == extension.upper(),
        )
    )
    other = User(
        login="export_other",
        role="TEACHER",
        first_name="Другой",
        last_name="Педагог",
        password_hash="unused",
    )
    db.add(other)
    await db.flush()
    lesson.teacher_id = other.id
    await db.flush()
    assert (await client.get(path)).status_code == 404


async def test_certificate_requires_finished_lesson_and_enrolled_student(
    db: AsyncSession, client: AsyncClient
) -> None:
    lesson, rows = await lesson_fixture(db, 1)
    rows[0].delivered_at = datetime.now(UTC)
    await sign_in(client, "teacher", "teacher")
    path = f"/api/v1/teacher/lessons/{lesson.id}/students/{rows[0].student_id}/certificate.pdf"
    assert (await client.get(path)).status_code == 409
    lesson.status, lesson.finished_at = "FINISHED", datetime(2026, 9, 27, 10, tzinfo=UTC)
    await db.flush()
    assert (
        await client.get(path.replace(str(rows[0].student_id), str(uuid4())))
    ).status_code == 404
    response = await client.get(path)
    assert response.status_code == 200 and response.content.startswith(b"%PDF-")
    text = "".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert "27.09.2026 13:00 МСК" in text and "Выдано карточек: 1" in text
    assert await db.scalar(
        select(AuditLog.id).where(AuditLog.action == "TRAINING_CERTIFICATE_EXPORTED")
    )


def test_xml_configuration_round_trip() -> None:
    value = RuntimeConfiguration()
    value.logging.http_access = False
    value.database.api_pool_size = 12
    assert configuration_xml.decode(configuration_xml.encode(value)) == value


@pytest.mark.parametrize(
    "kind",
    [
        "dtd",
        "duplicate",
        "missing",
        "unknown",
        "nested",
        "boolean",
        "range",
        "schema",
        "oversized",
        "encoding",
    ],
)
def test_xml_configuration_rejects_unsafe_ambiguous_and_invalid_values(kind: str) -> None:
    raw = configuration_xml.encode(RuntimeConfiguration())
    if kind == "dtd":
        raw = b'<!DOCTYPE r [<!ENTITY x SYSTEM "file:///etc/passwd">]>' + raw
    elif kind == "duplicate":
        raw = raw.replace(b"</logging>", b"<level>INFO</level></logging>")
    elif kind == "missing":
        raw = raw.replace(b"<level>INFO</level>", b"")
    elif kind == "unknown":
        raw = raw.replace(b"level", b"unknown")
    elif kind == "nested":
        raw = raw.replace(b">INFO<", b"><nested/>INFO<")
    elif kind == "boolean":
        raw = raw.replace(b">true<", b">yes<")
    elif kind == "range":
        raw = raw.replace(b">20<", b">9999<")
    elif kind == "schema":
        raw = raw.replace(b'schema="1"', b'schema="2"')
    elif kind == "oversized":
        raw += b" " * 65536
    else:
        raw = b"\xff"
    with pytest.raises((ValueError, ParseError)):
        configuration_xml.decode(raw)


async def test_import_only_prepares_draft_and_requires_admin(
    db: AsyncSession, client: AsyncClient
) -> None:
    path = "/api/v1/admin/operations/configuration/import.xml"
    value = RuntimeConfiguration()
    value.database.api_pool_size = 12
    content = configuration_xml.encode(value)
    await sign_in(client, "teacher", "teacher")
    assert (await client.post(path, content=content)).status_code == 403
    await client.post("/api/v1/auth/logout")
    await sign_in(client, "admin", "admin")
    before = await read_snapshot(db)
    response = await client.post(path, content=content)
    assert response.status_code == 200 and response.json()["applied"] is False
    assert response.json()["configuration"]["database"]["api_pool_size"] == 12
    assert await read_snapshot(db) == before
    assert (await client.post(path, content=b" " * 65537)).status_code == 413
    assert (await client.post(path, content=b"<broken>")).status_code == 400
