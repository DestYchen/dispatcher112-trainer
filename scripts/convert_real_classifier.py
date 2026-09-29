"""Convert the customer's real incident classifier (ЕКП, v046_24) into the trainer's import format.

Source: one sheet, header in rows 1-3. Columns A-D build the numeric code, F holds group names on
group header rows, G/H/I are the three levels of formal tags the 112 operator clicks, K is the final
incident type, and columns N.. are notification lists: a non-empty cell means the service receives
the card. Some services have several columns that only fire when a modifier tag is chosen
(e.g. "Служба 101 (выбран признак НД - НЕТ ДОСТУПА)"); those become modifier_services.

Output: data/classifier.xlsx with sheets services / incident_groups / incident_types /
modifier_services, which app.seeds.import_classifier reads unchanged.

Usage: python scripts/convert_real_classifier.py <source.xlsx> [output.xlsx]
"""

import json
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook

HIDDEN_FOR_OPERATOR = "Не отображается оператору 112"

# code, display name, visible on the card, base columns, {modifier: columns that fire only with it}.
# Hidden services receive cards for statistics/analytics only; the АРМ-112 memo says such services
# are not drawn in the notification list.
SERVICES: list[tuple[str, str, bool, list[int], dict[str, list[int]]]] = [
    ("S101", "Служба 101 (МЧС)", True, [13], {"NO_ACCESS": [14]}),
    ("ODS_PSC", "ОДС ПСЦ", True, [15], {"THREAT_TO_PEOPLE": [16], "VICTIMS": [17], "NO_ACCESS": [18]}),
    ("MGPSS", "МГПСС", True, [19], {}),
    ("S102", "Служба 102 (МВД)", True, [20], {"VICTIMS": [22], "FATALITIES": [22]}),
    ("S103", "Служба 103 (СМП)", True, [23], {"VICTIMS": [24]}),
    ("S104", "Служба 104 (Мосгаз)", True, [26], {}),
    ("CEMP", "ЦЭМП", True, [28], {"THREAT_TO_PEOPLE": [29], "VICTIMS": [30, 31], "FATALITIES": [30]}),
    ("FSB", "ФСБ", True, [33], {}),
    ("MOSOBLGAZ", "Мособлгаз", True, [35], {}),
    ("AVTODOR", "Автомобильные дороги", True, [36, 76], {}),
    ("MOSGORTRANS", "Мосгортранс", True, [37], {"ROAD_BLOCKED": [39]}),
    ("GORHOZ", "Городское хозяйство", True, [40], {}),
    ("GORMOST", "ГБУ «Гормост»", True, [41, 42, 43, 44], {}),
    ("KANAL", "Канал имени Москвы", True, [45], {}),
    ("MGTS", "МГТС", True, [46, 47], {}),
    ("METRO", "Метрополитен", True, [48], {}),
    ("MOSVODOKANAL", "Мосводоканал", True, [49], {}),
    ("MOEK", "МОЭК", True, [50], {}),
    ("ROSSETI", "Россети Московский регион", True, [51], {}),
    ("OEK", "ОЭК", True, [52], {}),
    ("MOSLIFT", "Мослифт", True, [53], {}),
    ("CODD", "ЦОДД", True, [54], {}),
    ("DEP_ZHKH", "Департамент ЖКХ", True, [55], {}),
    ("MOSKOLLEKTOR", "Москоллектор", True, [59], {}),
    ("RZD", "РЖД (Московская ЖД)", True, [60], {}),
    ("VODHOZ", "Центррегионводхоз", True, [62], {}),
    ("OATI", "ОАТИ", True, [64], {}),
    ("MOSVODOSTOK", "Мосводосток", True, [65], {}),
    ("DPPOOS", "Департамент ППиООС", True, [66], {}),
    ("RITUAL", "ДТУ (Ритуал)", True, [71], {}),
    ("DTU", "ДТУ", True, [72], {}),
    ("ROSGVARDIA", "Росгвардия", True, [73], {}),
    ("DDS_DISTRICT", "ДДС района (территориальный ОИВ)", True, [74], {}),
    ("DDS_TINAO", "ДДС ТиНАО (территориальный ОИВ)", True, [75], {}),
    ("DEP_STROY", "Департамент строительства", True, [77, 78], {}),
    ("MZHI", "Мосжилинспекция", True, [80], {}),
    ("ORG_PEREVOZOK", "Организатор перевозок", True, [92], {"ROAD_BLOCKED": [93]}),
    ("CITYENERGO", "Ситиэнерго", True, [97], {}),
    ("DEP_GRSTROY", "Департамент гражданского строительства", True, [98], {}),
    # Information-only recipients.
    ("RBIPK", "Департамент РБиПК", False, [56, 57], {}),
    ("MAYOR", "Аппарат Мэра", False, [58], {}),
    ("DEP_OBR", "Департамент образования", False, [61], {}),
    ("KOMENDATURA", "Военная комендатура", False, [63], {}),
    ("TSZN", "Департамент ТСЗН", False, [67], {}),
    ("RSVO", "РСВО", False, [68], {}),
    ("EVAZHD", "ЭВАЖД", False, [69], {}),
    ("MSPPN", "МСППН", False, [70], {}),
    ("VET", "Комитет ветеринарии", False, [79], {}),
    ("CULTURE", "Департамент культуры", False, [81], {}),
    ("CSA", "ГКУ ЦСА им. Глинки", False, [82], {}),
    ("NTU", "ГКУ НТУ", False, [83], {}),
    ("FSO", "ФСО", False, [84], {}),
    ("GUP_MSR", "ГУП МСР", False, [85, 86], {}),
    ("TOURISM", "Комитет по туризму", False, [87], {}),
    ("DGP", "Департамент градостроительной политики", False, [88, 89], {}),
    ("CUKB", "ЦУКБ Минобороны", False, [90, 91], {}),
    ("ECOMON", "Мосэкомониторинг", False, [94], {}),
    ("MO_RHBZ", "Минобороны РХБЗ", False, [95, 96], {}),
]


def cell(row: tuple, index: int) -> str:
    value = row[index] if index < len(row) else None
    return "" if value is None else str(value).strip()


def difficulty(levels: int, services: int, group_name: str) -> int:
    # Prior only: more tag levels and a longer notification list mean more to get right.
    # Elo on real attempts replaces this once students have played the type.
    hard_groups = ("Взрыв", "Угрозы", "Аварии на опасных", "Обрушения", "террорист")
    bonus = 1 if any(word in group_name for word in hard_groups) else 0
    return max(1, min(10, 1 + levels + services // 3 + bonus))


def convert(source: Path, target: Path) -> dict[str, int]:
    rows = list(load_workbook(source, read_only=True, data_only=False).worksheets[0].iter_rows(values_only=True))
    groups: list[tuple[str, str]] = []
    types: list[list] = []
    seen_attrs: dict[str, str] = {}
    duplicates = 0
    group_code = ""
    group_name = ""
    for row in rows[3:]:
        final = cell(row, 10)
        if not final:
            if cell(row, 4).isdigit() and cell(row, 5):
                group_code, group_name = f"{int(cell(row, 4)):02}", cell(row, 5)
                groups.append((group_code, group_name))
            continue
        level1 = cell(row, 6)
        if not group_code or level1 == HIDDEN_FOR_OPERATOR:
            continue
        parts = [cell(row, i) for i in range(4)]
        if not all(p.isdigit() for p in parts):
            continue
        code = str(int(parts[0]) * 1_000_000 + int(parts[1]) * 10_000 + int(parts[2]) * 100 + int(parts[3]))
        attrs = {k: v for k, v in (("level1", level1), ("level2", cell(row, 7)), ("level3", cell(row, 8))) if v}
        key = json.dumps(attrs, ensure_ascii=False, sort_keys=True)
        if key in seen_attrs:
            # Same clicked tags must give one type; keep the first row, like the memo's rule.
            duplicates += 1
            continue
        seen_attrs[key] = code
        services = [code_ for code_, _, _, base, _ in SERVICES if any(cell(row, i) for i in base)]
        visible = sum(1 for code_, _, shown, _, _ in SERVICES if shown and code_ in services)
        types.append([code, group_code, final[:500], key, difficulty(len(attrs), visible, group_name),
                      ",".join(services)])

    modifiers: dict[str, set[str]] = defaultdict(set)
    for code_, _, _, _, conditional in SERVICES:
        for modifier in conditional:
            modifiers[modifier].add(code_)

    workbook = Workbook()
    readme = workbook.active
    readme.title = "README"
    readme.append([f"Сконвертировано из {source.name} скриптом scripts/convert_real_classifier.py."])
    readme.append(["Признаки 1-3 = уровни опросной карты 112; службы = непустые ячейки списка оповещения."])
    sheet = workbook.create_sheet("services")
    sheet.append(["code", "name", "short_name", "is_visible", "sort_order"])
    for i, (code_, name, visible, _, _) in enumerate(SERVICES):
        sheet.append([code_, name, name.split(" (")[0][:64], visible, i])
    sheet = workbook.create_sheet("incident_groups")
    sheet.append(["code", "name", "sort_order"])
    for i, (code_, name) in enumerate(groups):
        sheet.append([code_, name[:255], i])
    sheet = workbook.create_sheet("incident_types")
    sheet.append(["code", "group_code", "name", "attributes", "difficulty", "services"])
    for row in types:
        sheet.append(row)
    sheet = workbook.create_sheet("modifier_services")
    sheet.append(["modifier", "service_code"])
    for modifier, codes in sorted(modifiers.items()):
        for code_ in sorted(codes):
            sheet.append([modifier, code_])
    workbook.save(target)
    return {"groups": len(groups), "types": len(types), "services": len(SERVICES),
            "skipped_duplicate_tags": duplicates, "types_without_services": sum(1 for t in types if not t[5])}


if __name__ == "__main__":
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) > 2 else Path(__file__).resolve().parents[1] / "data" / "classifier.xlsx"
    print(convert(src, dst))
