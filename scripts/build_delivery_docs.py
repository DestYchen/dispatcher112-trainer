"""Create local DOCX/PDF manuals and a PPTX/PDF presentation from project documents."""

import argparse
import html
import re
import subprocess
from pathlib import Path

from docx import Document
from docx.shared import Pt as WordPt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

import compose

ROOT = Path(__file__).resolve().parents[1]
MANUALS = {
    "user-guide": (
        "Руководство пользователя",
        ["README.md", "docs/TELEPHONY.md", "docs/LEARNING-AI.md"],
    ),
    "administrator-guide": (
        "Руководство администратора",
        [
            "docs/DEPLOY.md",
            "docs/TLS.md",
            "docs/LOCAL-OPERATIONS.md",
            "docs/RUNTIME-CONFIGURATION.md",
            "docs/FILE-INTEGRITY.md",
            "docs/RECOVERY.md",
            "docs/RECOVERY-SWITCH.md",
            "docs/UPDATES.md",
            "docs/CLUSTER.md",
        ],
    ),
    "technical-description": (
        "Техническое описание",
        ["docs/ARCHITECTURE.md", "API.md", "docs/LIBRARIES.md"],
    ),
    "acceptance-status": (
        "Состояние реализации и результаты приёмки",
        ["docs/DELIVERY.md", "PROGRESS.md"],
    ),
}
SLIDES = [
    (
        "АРМ-112 · учебный симулятор",
        [
            "Локальная подготовка диспетчеров ДДС",
            "Два процесса: заполнение карточки и действия по поступившей карточке",
            "Работающая локальная поставка и измеренные результаты приёмки",
        ],
        None,
    ),
    (
        "Рабочее место обучающегося",
        [
            "Очередь, карточка, действия и телефония",
            "Серверные сроки и обязательные комментарии",
            "Сохранение ввода и повтор действий после разрыва связи",
        ],
        "card-1440.png",
    ),
    (
        "Контроль учебного процесса",
        [
            "Преподаватель назначает и завершает занятия",
            "Пульт участников, наблюдение и обратная связь",
            "Воспроизводимая оценка и отдельное переопределение с аудитом",
        ],
        "teacher-live-1440.png",
    ),
    (
        "Учебные материалы и телефония",
        [
            "Группы, библиотека, модули и история обучения",
            "Локальная генерация черновиков с проверкой преподавателем",
            "SIP/WebRTC, голосовые реплики, записи и доклады",
        ],
        None,
    ),
    (
        "Администрирование",
        [
            "Политики доступа, параметры БД/SIP/журналирования",
            "Полные копии, подписанные обновления и проверка файлов",
            "Локальные операции сервисов, ресурсы и переключение восстановления",
        ],
        None,
    ),
    (
        "Локальная архитектура",
        [
            "React → HTTPS/NGINX → FastAPI → PostgreSQL",
            "Redis: очереди, общие события и присутствие",
            "Asterisk, LanguageTool и локальная модель работают внутри контура",
        ],
        None,
    ),
    (
        "Обмен и комплект поставки",
        [
            "Excel/CSV/XML/PDF отчёты, подтверждения обучения и MP3/WAV",
            "XML-параметры и методические материалы",
            "Исходный код, локальные образы, данные, инструкции и лицензии",
        ],
        None,
    ),
    (
        "Приёмка и внешние исходные материалы",
        [
            "Проверены восстановление, отказ backend и три браузера",
            "100 сессий: p95 677 мс; пять браузеров: карточка за 454–591 мс",
            "Для предметной сдачи нужны оригинальные билеты, классификаторы и целевой аппаратный стенд",
        ],
        None,
    ),
]


def tokens():
    return dict(
        re.findall(
            r"(--[\w-]+)\s*:\s*([^;]+);",
            (ROOT / "frontend/src/styles/tokens.css").read_text(encoding="utf-8"),
        )
    )


def plain(text):
    return (
        re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
        .replace("**", "")
        .replace("`", "")
    )


def blocks(paths):
    for relative in paths:
        yield "heading", relative, 1
        code = False
        for line in (ROOT / relative).read_text(encoding="utf-8").splitlines():
            if line.startswith("```"):
                code = not code
                continue
            if not line.strip():
                continue
            match = re.match(r"^(#{1,6})\s+(.+)$", line)
            if match and not code:
                yield "heading", plain(match[2]), min(len(match[1]) + 1, 4)
            elif code:
                yield "code", line, 0
            elif line.startswith("|") and re.fullmatch(r"[| :\-]+", line):
                continue
            else:
                yield "text", plain(line), 0


def render_pdf(content, path):
    code = (
        "import sys; from weasyprint import HTML; from app.domain.report_pdf import LocalFonts; "
        "sys.stdout.buffer.write(HTML(string=sys.stdin.buffer.read().decode('utf-8'), "
        "url_fetcher=LocalFonts()).write_pdf())"
    )
    result = subprocess.run(
        compose.command(
            ROOT, ["exec", "-T", "--index", "1", "backend", "python", "-c", code]
        ),
        input=content.encode("utf-8"),
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
    )
    if result.returncode or not result.stdout.startswith(b"%PDF-"):
        raise RuntimeError(
            "Local PDF generation failed; backend and local fonts are required"
        )
    path.write_bytes(result.stdout)


def document(title, paths, output, styles):
    word = Document()
    word.styles["Normal"].font.name = "PT Sans"
    word.styles["Normal"].font.size = WordPt(10)
    word.add_heading(title, 0)
    word.add_paragraph(
        "Учебный симулятор АРМ-112. Сформировано из документации текущего исходного кода. Результаты проверок и внешние остатки приведены в PROGRESS.md."
    )
    body = [f"<h1>{html.escape(title)}</h1>"]
    for kind, text, level in blocks(paths):
        if kind == "heading":
            word.add_heading(text, level)
            body.append(f"<h{level}>{html.escape(text)}</h{level}>")
        elif kind == "code":
            word.add_paragraph(text, style="No Spacing")
            body.append(f"<pre>{html.escape(text)}</pre>")
        else:
            word.add_paragraph(text)
            body.append(f"<p>{html.escape(text)}</p>")
    word.save(output.with_suffix(".docx"))
    content = (
        f'<!doctype html><html lang="ru"><meta charset="utf-8"><style>{styles}</style>'
        + "".join(body)
        + "</html>"
    )
    output.with_suffix(".html").write_text(content, encoding="utf-8")
    render_pdf(content, output.with_suffix(".pdf"))


def presentation(output, values, styles):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    pages = []
    for index, (title, points, screenshot) in enumerate(SLIDES, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(
            values["--surface"].lstrip("#")
        )
        heading = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.45), Inches(12.3), Inches(1.05)
        ).text_frame
        heading.text = title
        heading.paragraphs[0].font.size = Pt(32)
        heading.paragraphs[0].font.bold = True
        heading.paragraphs[0].font.name = "PT Sans"
        heading.paragraphs[0].font.color.rgb = RGBColor.from_string(
            values["--select"].lstrip("#")
        )
        image = ROOT / "artifacts/ui-review/chrome" / screenshot if screenshot else None
        has_image = bool(image and image.is_file())
        frame = slide.shapes.add_textbox(
            Inches(0.55), Inches(1.8), Inches(4.15 if has_image else 12), Inches(4.8)
        ).text_frame
        frame.word_wrap = True
        for number, text in enumerate(points):
            paragraph = frame.paragraphs[0] if number == 0 else frame.add_paragraph()
            paragraph.text = text
            paragraph.font.name = "PT Sans"
            paragraph.font.size = Pt(21)
            paragraph.space_after = Pt(24)
            paragraph.font.color.rgb = RGBColor.from_string(values["--ink"].lstrip("#"))
        if has_image:
            slide.shapes.add_picture(
                str(image), Inches(4.9), Inches(1.8), width=Inches(7.8)
            )
        footer = slide.shapes.add_textbox(
            Inches(0.55), Inches(7), Inches(12), Inches(0.3)
        ).text_frame
        footer.text = f"АРМ-112 · учебный режим                                    {index} / {len(SLIDES)}"
        footer.paragraphs[0].font.size = Pt(12)
        pages.append(
            f"<section><h1>{html.escape(title)}</h1><ul>"
            + "".join(f"<li>{html.escape(point)}</li>" for point in points)
            + "</ul></section>"
        )
    deck.save(output.with_suffix(".pptx"))
    content = (
        '<!doctype html><html lang="ru"><meta charset="utf-8"><style>'
        + styles
        + "@page {size: A4 landscape} section {break-after:page} li {margin-bottom:1em}</style>"
        + "".join(pages)
        + "</html>"
    )
    output.with_suffix(".html").write_text(content, encoding="utf-8")
    render_pdf(content, output.with_suffix(".pdf"))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output", type=Path, default=ROOT / "delivery/2026-09-27/documents"
    )
    args = parser.parse_args()
    output = args.output.resolve()
    if not output.is_relative_to(ROOT / "delivery"):
        raise ValueError("Delivery documents must stay within delivery/")
    output.mkdir(parents=True, exist_ok=True)
    values = tokens()
    styles = (
        "@font-face{font-family:PTSans;src:url(file:///data/fonts/PTSans-400.ttf)}"
        "@font-face{font-family:PTSans;src:url(file:///data/fonts/PTSans-700.ttf);font-weight:700}"
        f"@page{{size:A4;margin:{values['--pdf-margin']}}}"
        f"body{{font-family:PTSans;font-size:{values['--pdf-text']};color:{values['--ink']}}}"
        "p,pre{overflow-wrap:anywhere;white-space:pre-wrap} h1,h2,h3{break-after:avoid}"
    )
    for name, (title, paths) in MANUALS.items():
        document(title, paths, output / name, styles)
        print(f"Created {name}.docx and .pdf", flush=True)
    presentation(output / "presentation", values, styles)
    print("Created presentation.pptx and .pdf", flush=True)


if __name__ == "__main__":
    main()
